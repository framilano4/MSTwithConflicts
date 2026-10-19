#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import argparse
import heapq
import os
import random
import sys
import time

# ==========================================================================
#          CARTELLA_ISTANZE = "istanze"
#          CARTELLA_ISTANZE = r"C:\Users\nome\Desktop\istanze"
# ==========================================================================
CARTELLA_ISTANZE = r"C:\Users\frank\OneDrive\Desktop\TESI\python\Istanze tipo 1"


# --------------------------------------------------------------------------
# Lettura istanza
# --------------------------------------------------------------------------
class Istanza:
    def __init__(self, nome, V, radice, archi, conflitti):
        self.nome = nome
        self.V = V
        self.radice = radice
        self.archi = archi          # lista di (i, j, peso), indice = id arco
        self.conflitti = conflitti  # lista di coppie (id_arco_a, id_arco_b)
        self.E = len(archi)
        # lista di adiacenza dei conflitti per arco
        self.conf = [[] for _ in range(self.E)]
        for a, b in conflitti:
            self.conf[a].append(b)
            self.conf[b].append(a)


def leggi_istanza(percorso):
    """Formato: righe '#' = commento; poi 'V E K', radice, E archi 'i j peso',
    K conflitti 'i1 j1 i2 j2'."""
    token = []
    with open(percorso, "r", encoding="utf-8") as f:
        for riga in f:
            riga = riga.strip()
            if not riga or riga.startswith("#"):
                continue
            token.extend(riga.split())
    pos = 0
    V, E, K = int(token[0]), int(token[1]), int(token[2])
    radice = int(token[3])
    pos = 4

    archi = []
    indice = {}
    for e in range(E):
        i, j, w = int(token[pos]), int(token[pos + 1]), int(token[pos + 2])
        pos += 3
        archi.append((i, j, w))
        indice[(min(i, j), max(i, j))] = e

    conflitti = []
    for _ in range(K):
        i1, j1, i2, j2 = (int(t) for t in token[pos:pos + 4])
        pos += 4
        a = indice[(min(i1, j1), max(i1, j1))]
        b = indice[(min(i2, j2), max(i2, j2))]
        if a != b:
            conflitti.append((a, b))

    return Istanza(os.path.basename(percorso), V, radice, archi, conflitti)


# --------------------------------------------------------------------------
# Componenti dell'euristica
# --------------------------------------------------------------------------
def costruisci_S(ist, ordine):
    """Scorre gli archi nell'ordine dato e aggiunge un arco a S solo se non
    e' in conflitto con nessun arco gia' inserito."""
    in_S = [False] * ist.E
    S = []
    conf = ist.conf
    for e in ordine:
        ok = True
        for f in conf[e]:
            if in_S[f]:
                ok = False
                break
        if ok:
            in_S[e] = True
            S.append(e)
    return S


def prim_su_S(ist, S):
    """Prim sul sottografo formato da S, partendo dalla radice.
    Restituisce (costo, lista archi) oppure None se il sottografo non e'
    connesso (iterazione da scartare)."""
    V = ist.V
    adj = [[] for _ in range(V)]
    for e in S:
        i, j, w = ist.archi[e]
        adj[i].append((w, j, e))
        adj[j].append((w, i, e))

    visitato = [False] * V
    visitato[ist.radice] = True
    n_vis = 1
    heap = list(adj[ist.radice])
    heapq.heapify(heap)
    costo = 0
    albero = []
    while heap and n_vis < V:
        w, v, e = heapq.heappop(heap)
        if visitato[v]:
            continue
        visitato[v] = True
        n_vis += 1
        costo += w
        albero.append(e)
        for t in adj[v]:
            if not visitato[t[1]]:
                heapq.heappush(heap, t)

    if n_vis < V:
        return None
    return costo, albero


def genera_ordine(ist, k, rng):
    """Ordine completamente casuale degli archi (diverso a ogni tentativo)."""
    ordine = list(range(ist.E))
    rng.shuffle(ordine)
    return ordine


def euristica_multistart(ist, n_start=5000, tempo_max=None, seed=0):
    rng = random.Random(seed)
    t0 = time.perf_counter()
    best_costo, best_albero = None, None
    ammissibili = 0
    tentativi = 0
    for k in range(n_start):
        if tempo_max is not None and time.perf_counter() - t0 > tempo_max:
            break
        tentativi += 1
        ordine = genera_ordine(ist, k, rng)
        S = costruisci_S(ist, ordine)
        ris = prim_su_S(ist, S)
        if ris is None:
            continue                     # sottografo non connesso: scarto
        ammissibili += 1
        costo, albero = ris
        if best_costo is None or costo < best_costo:
            best_costo, best_albero = costo, albero
    tempo = time.perf_counter() - t0
    return best_costo, best_albero, tempo, tentativi, ammissibili


# --------------------------------------------------------------------------
# Controllo di ammissibilita'
# --------------------------------------------------------------------------
def verifica_soluzione(ist, albero, costo_dichiarato):
    """Controlla: V-1 archi distinti, albero ricoprente (union-find), nessuna
    coppia in conflitto, costo ricalcolato = costo dichiarato."""
    if len(set(albero)) != ist.V - 1 or len(albero) != ist.V - 1:
        return False, "numero di archi diverso da V-1"
    padre = list(range(ist.V))

    def trova(x):
        while padre[x] != x:
            padre[x] = padre[padre[x]]
            x = padre[x]
        return x

    for e in albero:
        i, j, _ = ist.archi[e]
        a, b = trova(i), trova(j)
        if a == b:
            return False, "presente un ciclo"
        padre[a] = b
    insieme = set(albero)
    for a, b in ist.conflitti:
        if a in insieme and b in insieme:
            return False, "presente una coppia di archi in conflitto"
    if sum(ist.archi[e][2] for e in albero) != costo_dichiarato:
        return False, "costo non coerente"
    return True, "ok"


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def salva_excel(percorso, risultati):
    """Salva i risultati in un file Excel (.xlsx). Se openpyxl non e'
    disponibile, salva un CSV con lo stesso contenuto."""
    intestazioni = ["Istanza", "V", "E", "K", "Obiettivo", "Tempo (s)",
                    "Tentativi", "Tentativi ammissibili"]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Risultati"
        ws.append(intestazioni)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in risultati:
            ws.append(r)
        for riga in ws.iter_rows(min_row=2, min_col=6, max_col=6):
            riga[0].number_format = "0.000"
        for col, larg in zip("ABCDEFGH", (22, 6, 6, 6, 12, 12, 11, 22)):
            ws.column_dimensions[col].width = larg
        wb.save(percorso)
        return percorso
    except ImportError:
        import csv
        percorso = os.path.splitext(percorso)[0] + ".csv"
        with open(percorso, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(intestazioni)
            w.writerows(risultati)
        return percorso


def elenco_file(percorsi):
    out = []
    for p in percorsi:
        if os.path.isdir(p):
            for nome in sorted(os.listdir(p)):
                if nome.lower().endswith(".txt"):
                    out.append(os.path.join(p, nome))
        else:
            out.append(p)
    return out


def main():
    ap = argparse.ArgumentParser(description="Euristica multi-start per MSTC")
    ap.add_argument("istanze", nargs="*", help="file istanza o cartelle (default: CARTELLA_ISTANZE)")
    ap.add_argument("--starts", type=int, default=5000, help="n. massimo di ordinamenti")
    ap.add_argument("--tempo", type=float, default=None, help="tempo max euristica (s) per istanza")
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--output", default="risultati_mstc.xlsx", help="file Excel dei risultati")
    args = ap.parse_args()
    if not args.istanze:
        args.istanze = [CARTELLA_ISTANZE]

    righe = []
    risultati = []
    for percorso in elenco_file(args.istanze):
        ist = leggi_istanza(percorso)
        costo, albero, t, tent, amm = euristica_multistart(
            ist, args.starts, args.tempo, args.seed)

        if costo is None:
            print(f"{ist.nome}: nessuna soluzione ammissibile trovata "
                  f"({tent} tentativi, tempo {t:.3f}s)")
            righe.append((ist.nome, None, t, "-"))
            risultati.append([ist.nome, ist.V, ist.E, len(ist.conflitti), "n/d", round(t, 4), tent, amm])
            continue

        ok, msg = verifica_soluzione(ist, albero, costo)
        if not ok:
            print(f"ERRORE {ist.nome}: soluzione non valida ({msg})", file=sys.stderr)
            sys.exit(1)

        print(f"{ist.nome}: obiettivo = {costo}  tempo = {t:.3f} s  "
              f"(tentativi {tent}, ammissibili {amm})")
        righe.append((ist.nome, costo, t, f"{amm}/{tent}"))
        risultati.append([ist.nome, ist.V, ist.E, len(ist.conflitti), costo, round(t, 4), tent, amm])

    if len(righe) > 1:
        print("\n{:<30}{:>12}{:>12}{:>14}".format("istanza", "obiettivo", "tempo (s)", "ammiss/tent"))
        for nome, c, t, r in righe:
            print("{:<30}{:>12}{:>12.3f}{:>14}".format(nome, "-" if c is None else c, t, r))

    if risultati:
        salvato = salva_excel(args.output, risultati)
        print(f"\nRisultati salvati in: {os.path.abspath(salvato)}")


if __name__ == "__main__":
    main()
